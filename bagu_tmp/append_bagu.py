# -*- coding: utf-8 -*-
import re, shutil
from io import BytesIO
from docx import Document

SRC = r'C:\颜-个人\澳洲，新西兰，爱尔兰\Australia\国立\Resume\项目八股_完整版.docx'
BAK = r'D:\新建文件夹\java\javacode\nautix-frontend\bagu_tmp\项目八股_完整版.backup.docx'
TMP = r'D:\新建文件夹\java\javacode\nautix-frontend\bagu_tmp'

files = ['01_java.md', '02_mysql.md', '02_redis.md',
         '03_network.md', '04_os.md', '04_spring.md']

shutil.copyfile(SRC, BAK)
doc = Document(SRC)

def add_body(text):
    p = doc.add_paragraph()
    m = re.match(r'^(\u3010[^\u3011]+\u3011)(.*)$', text)
    if m:
        r = p.add_run(m.group(1)); r.bold = True
        p.add_run(m.group(2))
    else:
        p.add_run(text)

def flush(paras):
    for raw in paras:
        t = raw.strip()
        if not t:
            continue
        if t.startswith('#####'):
            doc.add_paragraph(t[5:].strip(), style='Heading 3')
        elif t.startswith('####'):
            doc.add_paragraph(t[4:].strip(), style='Heading 2')
        elif t.startswith('###'):
            doc.add_paragraph(t[3:].strip(), style='Heading 1')
        else:
            add_body(t)

for fn in files:
    with open(TMP + '\\' + fn, encoding='utf-8') as f:
        buf = []
        for line in f:
            s = line.rstrip('\n')
            if s.lstrip().startswith('#'):
                flush(buf); buf = []
                flush([s])
                continue
            if s.strip() == '':
                if buf:
                    flush(buf); buf = []
                continue
            buf.append(s)
        flush(buf)

out = BytesIO()
doc.save(out)
merged = TMP + '\\' + '项目八股_完整版_含基础八股.docx'
with open(merged, 'wb') as f:
    f.write(out.getvalue())

import sys
coped = False
if '--try-src' in sys.argv:
    try:
        shutil.copyfile(merged, SRC)
        coped = True
    except PermissionError:
        print('SRC_COPY_BLOCKED')

chk = Document(SRC if coped else merged)
print('total paras:', len(chk.paragraphs))
