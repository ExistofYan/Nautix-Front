from docx import Document

M = r'D:\新建文件夹\java\javacode\nautix-frontend\bagu_tmp\项目八股_完整版_含基础八股.docx'
d = Document(M)

h1 = [p.text for p in d.paragraphs if p.style.name == 'Heading 1']
h2 = [p.text for p in d.paragraphs if p.style.name == 'Heading 2']
h3 = [p.text for p in d.paragraphs if p.style.name == 'Heading 3']
print('total paras:', len(d.paragraphs))
print('H1:', h1)
print('H2 count:', len(h2))
print('H3 count:', len(h3))
print('--- H3 list ---')
for t in h3:
    print(t)
