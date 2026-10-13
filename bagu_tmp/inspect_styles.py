from docx import Document
import json

path = r'C:\颜-个人\澳洲，新西兰，爱尔兰\Australia\国立\Resume\项目八股_完整版.docx'
doc = Document(path)

targets = ['一、DDEC', '五、Java 核心八股补全', 'HashMap / ConcurrentHashMap',
           'J1. HashMap 的底层结构', '【解释】', '【面试回答】', '铁律一', '六、考前突击']
found = {}
for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    for key in targets:
        if key not in found and t.startswith(key) or (key in ('【解释】', '【面试回答】') and t.startswith(key)):
            found[key] = {
                'style': p.style.name,
                'sample': t[:40],
                'runs_bold': [bool(r.bold) for r in p.runs[:3]],
            }

print(json.dumps(found, ensure_ascii=False, indent=1))
print('para count:', len(doc.paragraphs))
